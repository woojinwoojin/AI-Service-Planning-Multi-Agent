"""사용자 여정 API 통합 테스트 (로드맵 Phase 7).

입력 자동완성 → 생성 → 이력 재조회 → 사용자 수정(HITL) → 내보내기(DOCX/PPTX/MD/JSON)
까지를 **한 흐름**으로 검증한다. 개별 라우트 단위 테스트(test_routes.py 등)와 달리, 실제
사용자가 밟는 순서대로 상태가 응답·이력·산출물에 일관되게 흐르는지 본다.

⚠ 이름 주의(외부 리뷰 3차 D-5): 브라우저 E2E 가 아니라 **FastAPI TestClient 기반 API 통합**
테스트다(그래서 파일명이 test_api_journey). 프런트엔드(진행 단계 표시·수정 후 메타 반영·오류
메시지)는 여기서 커버되지 않는다 — 그쪽은 Playwright 등 브라우저 테스트가 필요하다.

- hermetic: 실제 LLM 미호출(dummy), 임시 DB·임시 OUTPUT_DIR — 반복·CI에서 안전.
- 실행 구조(serial/parallel)는 환경(.env WORKFLOW_MODE)에 의존하므로 값 자체는 단정하지 않는다.
- 외부 리뷰 P0~P2 수정(저장·수정 정합성)을 전 흐름 회귀로 고정한다.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from app.services import (
    docx_export,
    llm,
    markdown_export,
    pptx_export,
    reliability,
    store,
)

# quality_gate 가 판정하는 체크 항목(외부 리뷰 P1-3 로 contradicted_claims 추가됨)
_GATE_CHECKS = {"score", "critical_issues", "major_issues", "structure", "evidence", "contradicted_claims"}
# 저장·API 응답에 반드시 보존돼야 하는 PR-7/8 실행 기록(외부 리뷰 P0-1)
_PR78_KEYS = {"revision_strategy", "revised_section_ids", "revision_fallback_reason",
              "polish_applied", "polish_skip_reason"}


@pytest.fixture
def client(tmp_path, monkeypatch):
    """임시 DB + dummy LLM 로 관통. 실제 키·비용·네트워크 불필요."""
    monkeypatch.setattr(store, "DB_PATH", tmp_path / "projects.db")
    monkeypatch.setattr(llm, "is_dummy", lambda: True)
    return TestClient(app)


@pytest.fixture
def outputs_tmp(tmp_path, monkeypatch):
    """/run/save 가 파일을 쓰는 3개 OUTPUT_DIR 를 임시 경로로 돌려 저장소 오염 방지."""
    out = tmp_path / "outputs"
    monkeypatch.setattr(markdown_export, "OUTPUT_DIR", out)
    monkeypatch.setattr(docx_export, "OUTPUT_DIR", out)
    monkeypatch.setattr(pptx_export, "OUTPUT_DIR", out)
    return out


def test_e2e_idea_to_download(client):
    """아이디어 한 줄 → 자동완성 → 생성 → 이력 재조회 → DOCX/PPTX 내려받기 전 여정."""
    # 1) 입력 자동완성: 프로젝트명만으로 빈 항목 초안을 받는다.
    sug = client.post("/suggest", json={"project_name": "AI 반려식물 케어"}).json()
    assert {"description", "target_user", "problem", "keywords"} <= set(sug)

    # 2) 생성: 자동완성 결과를 그대로 입력으로 실행한다.
    run = client.post("/run", json={
        "project_name": "AI 반려식물 케어",
        "description": sug["description"],
        "target_user": sug["target_user"],
        "problem": sug["problem"],
        "keywords": sug["keywords"] if isinstance(sug["keywords"], list) else [],
    }).json()
    pid = run["project_id"]
    assert pid > 0
    assert run["final_draft"]                                  # 최종 기획서 본문 존재
    assert run["final_review_result"]                          # 최종본 재평가 포함
    assert run["verification_result"]                          # 근거 검증 수행
    assert run["workflow_mode"] in ("serial", "parallel")      # 실행 구조 태깅(값은 env 의존)
    assert "wall_time_ms" in run["usage"]                      # 관측치(지연)

    # 품질 게이트가 6개 체크(반대 근거 포함) 전부와 함께 표면화된다(P1-3).
    assert set(run["quality_gate"]["checks"]) == _GATE_CHECKS
    assert isinstance(run["quality_gate"]["release_ready"], bool)

    # PR-7/8 실행 기록이 API 응답에 담긴다(P0-1: 응답 측).
    assert _PR78_KEYS <= set(run)

    # 3) 이력 재조회: 저장 후 다시 열어도 실행 기록이 왜곡 없이 복원된다(P0-1: 저장 측).
    detail = client.get(f"/projects/{pid}").json()
    state = detail["state"]
    assert state["final_draft"] == run["final_draft"]          # 본문 왕복 동일
    assert state["revision_strategy"] == run["revision_strategy"]
    assert state["polish_applied"] == run["polish_applied"]
    assert state["quality_gate"]["checks"].keys() == run["quality_gate"]["checks"].keys()
    assert state.get("verification_result")                     # 검증 결과 보존
    ids = [p["id"] for p in client.get("/projects").json()["projects"]]
    assert pid in ids                                          # 목록에 존재

    # 4) 다운로드: 화면의 최종 기획서를 그대로 DOCX·PPTX 로 내려받는다.
    docx_r = client.post("/export/docx", json={
        "project_name": "AI 반려식물 케어", "markdown": run["final_draft"]})
    assert docx_r.status_code == 200
    assert "wordprocessingml" in docx_r.headers["content-type"]
    assert docx_r.content[:2] == b"PK"                         # zip 컨테이너
    doc = Document(io.BytesIO(docx_r.content))                 # 유효한 문서로 열림
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert reliability._MARKER in doc_text                     # 내보내기에 한계 문구 부착됨

    pptx_r = client.post("/export/pptx", json={
        "project_name": "AI 반려식물 케어", "markdown": run["final_draft"]})
    assert pptx_r.status_code == 200
    assert "presentationml" in pptx_r.headers["content-type"]
    assert pptx_r.content[:2] == b"PK"
    Presentation(io.BytesIO(pptx_r.content))                   # 유효한 프레젠테이션으로 열림


def test_e2e_revise_updates_all_surfaces(client, monkeypatch):
    """사용자 수정(HITL) 후 문서뿐 아니라 검증·품질·실행상태가 함께 갱신된다(외부 리뷰 P0-2)."""
    from app.api import routes

    run = client.post("/run", json={"project_name": "수정흐름", "problem": "P"}).json()
    pid = run["project_id"]

    # 재작성기가 마커 있는 수정본을 내도록 대체(더미 fallback 배치 아티팩트 회피).
    monkeypatch.setattr(routes.draft_writer, "revise", lambda state: {
        "final_draft": "# 수정흐름 기획서\n분명히 바뀐 수정본", "revision_count": 1,
        "logs": ["[revise] 재작성 완료 (revision=1)"]})

    rev = client.post("/revise", json={
        "project_name": "수정흐름", "draft": run["final_draft"],
        "revision_request": "더 구체적으로", "project_id": pid,
    }).json()

    # 수정본 기준으로 재계산된 판정·품질·실행상태·근거가 응답에 함께 온다(옛 값 잔존 방지).
    assert rev["project_id"] == pid
    assert rev["verification_result"]
    assert set(rev["quality_gate"]["checks"]) == _GATE_CHECKS
    assert "run_status" in rev
    assert "timing" in rev
    assert "evidence_registry" in rev
    assert "revision_strategy" in rev and "polish_applied" in rev

    # 이력도 수정본 기준으로 갱신된다(재조회 시 옛 문서·옛 검증이 나오지 않음).
    saved = client.get(f"/projects/{pid}").json()["state"]
    assert "분명히 바뀐 수정본" in saved["final_draft"]
    assert saved.get("verification_result")
    assert saved.get("quality_gate")


def test_e2e_revise_with_new_model_updates_history_model(client):
    """수정 시 다른 모델을 고르면 이력 목록의 model 컬럼도 함께 갱신된다(외부 리뷰 P2-8)."""
    run = client.post("/run", json={
        "project_name": "모델교체", "problem": "P", "model": "gpt-4o-mini"}).json()
    pid = run["project_id"]
    assert _project_model(client, pid) == "gpt-4o-mini"

    client.post("/revise", json={
        "project_name": "모델교체", "draft": run["final_draft"],
        "revision_request": "정리", "project_id": pid, "model": "gpt-4o"})

    # 목록의 model 이 최초 실행 모델이 아니라 수정 시 사용한 모델로 바뀌어야 한다.
    assert _project_model(client, pid) == "gpt-4o"


def test_e2e_run_and_save_writes_all_artifacts(client, outputs_tmp):
    """/run/save 가 MD·JSON·DOCX·PPTX 산출물을 실제로 파일로 남긴다(내보내기 전 흐름).

    KOSENA 대응(체크포인트 3) 이후로는 **KOSENA 본문·발표자료·AI 활용 로그**가 더해진다 —
    제출 형식이 본문/발표/AI 로그로 나뉘기 때문(PDF p4). 기존 4종은 그대로 남아야 한다.
    """
    r = client.post("/run/save", json={"project_name": "저장물", "problem": "P"}).json()

    for key in ("saved_md", "saved_json", "saved_docx", "saved_pptx",
                "saved_kosena_md", "saved_kosena_docx", "saved_kosena_pptx", "saved_ai_log_md"):
        p = Path(r[key])
        assert p.exists() and p.stat().st_size > 0, key

    # MD 에는 한계 문구가 부착돼 있다(KOSENA 본문에도 동일하게).
    md = Path(r["saved_md"]).read_text(encoding="utf-8")
    assert reliability._MARKER in md
    assert reliability._MARKER in Path(r["saved_kosena_md"]).read_text(encoding="utf-8")

    # 실행 JSON 에는 PR-7/8 실행 기록이 보존된다(_RUN_KEYS 확장, P0-1).
    saved = json.loads(Path(r["saved_json"]).read_text(encoding="utf-8"))
    assert _PR78_KEYS <= set(saved)
    assert "final_draft" in saved and "quality_gate" in saved


def test_e2e_stream_run_to_download(client):
    """스트리밍 실행(SSE) 경로: start→node*→done 소비 후 done 결과로 DOCX 내려받기."""
    types, done = [], None
    with client.stream("POST", "/run/stream",
                       json={"project_name": "스트림저장", "problem": "P"}) as r:
        assert r.status_code == 200
        assert "text/event-stream" in r.headers["content-type"]
        for line in r.iter_lines():
            if not line or not line.startswith("data: "):
                continue
            ev = json.loads(line[6:])
            types.append(ev["type"])
            if ev["type"] == "done":
                done = ev["result"]
    # SSE 계약: 첫 이벤트 start, 마지막 done
    assert types[0] == "start" and types[-1] == "done"
    assert done and done["project_id"] > 0 and done["final_draft"]

    # 스트리밍 결과의 최종 기획서를 그대로 다운로드로 이어간다.
    docx_r = client.post("/export/docx", json={
        "project_name": "스트림저장", "markdown": done["final_draft"]})
    assert docx_r.status_code == 200 and docx_r.content[:2] == b"PK"

    # 스트리밍 실행도 이력에 저장되어 재조회된다.
    assert client.get(f"/projects/{done['project_id']}").status_code == 200


# ── 외부 리뷰 3차 트랙 D: 응답 모델화 · API 하드닝 · SSE 하트비트 ─────────────────

def test_revise_returns_full_run_result(client, monkeypatch):
    """D-1: /revise 응답이 /run 과 같은 RunResult — 새 State 필드가 수정 응답에서 누락되지 않는다."""
    from app.api import routes

    run = client.post("/run", json={"project_name": "응답모델", "problem": "P"}).json()
    monkeypatch.setattr(routes.draft_writer, "revise", lambda state: {
        "final_draft": "# 응답모델 기획서\n수정본", "revision_count": 1, "logs": ["[revise] ok"]})
    rev = client.post("/revise", json={
        "project_name": "응답모델", "draft": run["final_draft"],
        "revision_request": "정리", "project_id": run["project_id"]})
    assert rev.status_code == 200
    body = rev.json()
    # /run 응답 키 전체가 그대로 온다(이전 수동 dict 에서 빠져 있던 메타 포함).
    assert set(body) == set(run)
    for key in ("revision_strategy", "revised_section_ids", "revision_fallback_reason",
                "polish_applied", "polish_skip_reason", "best_version",
                "reverted_from_revision", "failed_nodes", "state_version"):
        assert key in body, key
    assert body["state_version"] == run["state_version"]
    assert "수정본" in body["final_draft"]


def test_revise_unknown_project_id_returns_404(client):
    """D-4: 지정한 project_id 가 없으면 조용히 신규 저장하지 않고 404 (이력 쪼개짐 방지)."""
    r = client.post("/revise", json={
        "project_name": "없는프로젝트", "draft": "# x", "revision_request": "정리",
        "project_id": 999999})
    assert r.status_code == 404
    assert r.json()["error"]["code"] == "not_found"
    assert client.get("/projects").json()["projects"] == []     # 신규 레코드가 생기지 않았다


def test_projects_limit_is_bounded(client):
    """D-4: /projects limit 은 1~100 범위. 벗어나면 422(통일 오류 형식)."""
    client.post("/run", json={"project_name": "경계", "problem": "P"})
    assert client.get("/projects?limit=1").status_code == 200
    assert client.get("/projects?limit=100").status_code == 200
    for bad in (0, -1, 101):
        r = client.get(f"/projects?limit={bad}")
        assert r.status_code == 422, bad
        assert r.json()["error"]["code"] == "validation_error"


def test_sse_stream_emits_heartbeat_when_idle():
    """D-4: 이벤트 공백 구간에 SSE comment 를 흘려 reverse proxy 유휴 타임아웃을 피한다."""
    import time

    from app.api import routes

    def slow_events():
        time.sleep(0.25)                       # 노드 하나가 오래 걸리는 상황
        yield {"type": "node", "node": "draft", "order": 1}

    out = list(routes._with_heartbeat(slow_events(), heartbeat_sec=0.05))
    assert None in out                          # 하트비트 신호가 최소 1회
    assert out[-1] == {"type": "node", "node": "draft", "order": 1}   # 이벤트도 정상 전달


def test_sse_heartbeat_propagates_worker_error():
    """하트비트 래퍼가 워커 예외를 소비자 스레드로 다시 던져 기존 error 이벤트 경로를 유지한다."""
    from app.api import routes

    def boom_events():
        yield {"type": "start", "workflow_mode": "serial"}
        raise RuntimeError("실행 중 폭발")

    got = []
    with pytest.raises(RuntimeError, match="실행 중 폭발"):
        for ev in routes._with_heartbeat(boom_events(), heartbeat_sec=0.05):
            got.append(ev)
    assert got and got[0]["type"] == "start"


def test_sse_stream_tolerates_comment_frames(client):
    """하트비트 comment 가 섞여도 SSE 소비자는 start→node*→done 계약을 그대로 본다."""
    types, done = [], None
    with client.stream("POST", "/run/stream",
                       json={"project_name": "하트비트", "problem": "P"}) as r:
        assert r.status_code == 200
        for line in r.iter_lines():
            if not line or line.startswith(":"):    # comment 는 무시(규격)
                continue
            if not line.startswith("data: "):
                continue
            ev = json.loads(line[6:])
            types.append(ev["type"])
            if ev["type"] == "done":
                done = ev["result"]
    assert types[0] == "start" and types[-1] == "done"
    assert done["project_id"] > 0


# ── helpers ──────────────────────────────────────────────────────────────────

def _project_model(client: TestClient, pid: int) -> str:
    for p in client.get("/projects").json()["projects"]:
        if p["id"] == pid:
            return p["model"]
    raise AssertionError(f"project {pid} not in history")


def _only_md(out_dir) -> str:
    names = [p.name for p in out_dir.glob("*.md")]
    assert len(names) == 1, f"expected 1 md, got {names}"
    return names[0]
